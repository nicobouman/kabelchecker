"""Rekenkern Kabelchecker v4.

Leest regeltemplate_v4 (Excel) en toetst tracédelen in vier vaste stappen:
feiten afleiden -> vergunningplicht (2.15) -> regimes -> zorgplichtregels.
Elke toets heeft drie uitkomsten: waar, onwaar of onbekend (gegeven ontbreekt).

Gebruik vanaf de opdrachtregel (testgevallen draaien en terugschrijven):
    python kabelchecker_v4_engine.py regeltemplate_v4_AGV.xlsx --test
"""
from __future__ import annotations

import ast
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# ---------------------------------------------------------------- drie-waardige logica
class _Onbekend:
    _inst = None

    def __new__(cls):
        if cls._inst is None:
            cls._inst = super().__new__(cls)
        return cls._inst

    def __repr__(self):
        return "onbekend"

    def __bool__(self):
        raise TypeError("onbekend heeft geen waarheidswaarde")


ONBEKEND = _Onbekend()


def is_onbekend(v) -> bool:
    return v is ONBEKEND


class ExpressieFout(Exception):
    pass


def _vertaal(expr: str) -> str:
    s = re.sub(r"\bAND\b", " and ", expr)
    s = re.sub(r"\bOR\b", " or ", s)
    s = re.sub(r"\bNOT\b", " not ", s)
    return s.strip()


_CACHE: dict[str, ast.AST] = {}


def parse(expr: str) -> ast.AST:
    if expr not in _CACHE:
        try:
            _CACHE[expr] = ast.parse(_vertaal(expr), mode="eval").body
        except SyntaxError as e:
            raise ExpressieFout(f"Kan expressie niet lezen: {expr!r} ({e.msg})") from None
    return _CACHE[expr]


def _norm(v):
    if isinstance(v, str):
        return v.strip().lower()
    return v


class Context:
    """Zoekt namen op: invoer, geo-feiten, parameters en (recursief) begrippen."""

    def __init__(self, invoer: dict, geo: dict, params: dict, begrippen: dict[str, str]):
        self.invoer = {k: v for k, v in invoer.items() if v not in (None, "")}
        self.geo = geo
        self.params = params
        self.begrippen = begrippen
        self._memo: dict[str, Any] = {}
        self.gebruikt_onbekend: set[str] = set()

    def waarde(self, naam: str):
        if naam in self._memo:
            return self._memo[naam]
        if naam in self.params:
            v = self.params[naam]
        elif naam in self.invoer:
            v = self.invoer[naam]
        elif naam in self.geo:
            v = self.geo[naam]
            if v is None:
                v = ONBEKEND
        elif naam in self.begrippen and not self.begrippen[naam].startswith("geo"):
            self._memo[naam] = ONBEKEND  # bescherming tegen kringverwijzing
            v = evalueer(self.begrippen[naam], self)
        else:
            v = ONBEKEND
        if is_onbekend(v) and naam not in self.begrippen and naam not in self.params:
            self.gebruikt_onbekend.add(naam)
        self._memo[naam] = v
        return v


def _vergelijk(op, a, b):
    if is_onbekend(a) or is_onbekend(b):
        return ONBEKEND
    if isinstance(op, ast.In):
        return _norm(a) in [_norm(x) for x in b]
    if isinstance(op, ast.NotIn):
        return _norm(a) not in [_norm(x) for x in b]
    a, b = _norm(a), _norm(b)
    try:
        if isinstance(op, ast.Eq):
            return a == b
        if isinstance(op, ast.NotEq):
            return a != b
        a, b = float(a), float(b)
        return {ast.Lt: a < b, ast.LtE: a <= b, ast.Gt: a > b, ast.GtE: a >= b}[type(op)]
    except (TypeError, ValueError, KeyError):
        raise ExpressieFout(f"Kan {a!r} en {b!r} niet vergelijken")


def _ev(node, ctx: Context):
    if isinstance(node, ast.BoolOp):
        waarden = [_ev(v, ctx) for v in node.values]
        if isinstance(node.op, ast.And):
            if any(v is False for v in waarden):
                return False
            return ONBEKEND if any(is_onbekend(v) for v in waarden) else True
        if any(v is True for v in waarden):
            return True
        return ONBEKEND if any(is_onbekend(v) for v in waarden) else False
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.Not):
        v = _ev(node.operand, ctx)
        return ONBEKEND if is_onbekend(v) else (not _als_bool(v))
    if isinstance(node, ast.Compare):
        links = _ev(node.left, ctx)
        uitkomst: Any = True
        for op, rechts_node in zip(node.ops, node.comparators):
            rechts = _ev(rechts_node, ctx)
            r = _vergelijk(op, links, rechts)
            if r is False:
                return False
            if is_onbekend(r):
                uitkomst = ONBEKEND
            links = rechts
        return uitkomst
    if isinstance(node, ast.Name):
        if node.id in ("True", "False"):
            return node.id == "True"
        return ctx.waarde(node.id)
    if isinstance(node, ast.Constant):
        return node.value
    if isinstance(node, (ast.Tuple, ast.List)):
        return [_ev(e, ctx) for e in node.elts]
    raise ExpressieFout(f"Niet toegestaan in expressie: {ast.dump(node)[:60]}")


def _als_bool(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        return v.strip().lower() in ("ja", "true", "waar", "1")
    return bool(v)


def evalueer(expr: str, ctx: Context):
    """Geeft True, False of ONBEKEND. Lege expressie = True."""
    if not expr or not str(expr).strip():
        return True
    v = _ev(parse(str(expr)), ctx)
    return v if is_onbekend(v) else _als_bool(v)


# ---------------------------------------------------------------- template
@dataclass
class Template:
    params: dict[str, Any]
    begrippen: dict[str, str]
    zones: list[str]
    invoervelden: list[dict]
    vergunning: list[dict]
    regimes: list[dict]
    zorgplicht: list[dict]
    pad: Path | None = None


def _rijen(ws):
    it = ws.iter_rows(values_only=True)
    koppen = [str(k).strip() if k is not None else "" for k in next(it)]
    for r in it:
        if r is None or all(v in (None, "") for v in r):
            continue
        yield {k: v for k, v in zip(koppen, r) if k}


def laad_template(pad) -> Template:
    wb = load_workbook(pad, data_only=True)
    params = {}
    for r in _rijen(wb["Parameters"]):
        if r.get("naam"):
            params[str(r["naam"]).strip()] = float(r["waarde"])
    zones = [str(r["code"]).strip() for r in _rijen(wb["Zones"])
             if r.get("code") and re.fullmatch(r"[A-Z]+-[A-Z]+", str(r["code"]).strip())]
    begrippen = {str(r["naam"]).strip(): str(r.get("formule") or "").strip()
                 for r in _rijen(wb["Begrippen"]) if r.get("naam")}
    invoervelden = [r for r in _rijen(wb["Invoervelden"]) if r.get("veld")]

    def zones_van(r):
        return {z for z in zones if str(r.get(z, "")).strip().lower() == "ja"}

    vergunning = [dict(id=r["id"], lid=str(r["lid"]), zones=zones_van(r), voorwaarde=str(r.get("voorwaarde") or ""),
                       tekst=r.get("wsv-tekst") or "") for r in _rijen(wb["Vergunning_2.15"]) if r.get("id")]
    regimes = [dict(code=str(r["code"]).strip(), artikel=str(r.get("artikel") or ""), voorwaarde=str(r.get("voorwaarde") or ""),
                    omschrijving=r.get("omschrijving") or "") for r in _rijen(wb["Regimes"]) if r.get("code")]
    zorg = []
    for r in _rijen(wb["Zorgplicht"]):
        if not r.get("artikel"):
            continue
        zorg.append(dict(artikel=str(r["artikel"]).strip(), regime=str(r.get("regime") or "").strip(), zones=zones_van(r),
                         van_toepassing=str(r.get("van toepassing als") or ""), toets=str(r.get("toets") or ""),
                         soort=str(r.get("soort") or "voorschrift").strip().lower(),
                         voorschrift=r.get("voorschrift (bij voldoen)") or "", niet=r.get("tekst bij niet voldoen") or "",
                         wsv=r.get("wsv-tekst") or ""))
    return Template(params, begrippen, zones, invoervelden, vergunning, regimes, zorg, Path(pad))


# ---------------------------------------------------------------- toetsing van één deel
OORDELEN = ["vergunning", "maatwerk", "onvolledig", "aandacht", "voorschriften",
            "geen_specifieke_regels", "buiten_beperkingengebied"]


@dataclass
class Regeluitkomst:
    stap: str            # vergunning / regime / zorgplicht
    artikel: str
    uitkomst: str        # vergunning, voldoet, voldoet_niet, onbekend, voorschrift, aandacht, n.v.t.
    tekst: str = ""
    ontbrekend: list[str] = field(default_factory=list)


@dataclass
class Deeluitkomst:
    oordeel: str
    zones: list[str]
    regimes: list[str]
    regels: list[Regeluitkomst]
    vragen: list[str]    # velden om aan de vergunningverlener te vragen

    def artikelen(self, *uitkomsten):
        return [r.artikel for r in self.regels if r.uitkomst in uitkomsten]


def toets_deel(t: Template, invoer: dict, zones: set[str], geo: dict) -> Deeluitkomst:
    zones = {z for z in zones if z in t.zones}
    regels: list[Regeluitkomst] = []
    vragen: set[str] = set()
    if not zones:
        return Deeluitkomst("buiten_beperkingengebied", [], [], [], [])

    def ctx():
        return Context(invoer, geo, t.params, t.begrippen)

    # stap 2: vergunning
    for v in t.vergunning:
        if not (v["zones"] & zones):
            continue
        c = ctx()
        r = evalueer(v["voorwaarde"], c)
        if r is True:
            regels.append(Regeluitkomst("vergunning", v["lid"], "vergunning", v["tekst"]))
        elif is_onbekend(r):
            ontbr = sorted(c.gebruikt_onbekend)
            vragen |= set(ontbr)
            regels.append(Regeluitkomst("vergunning", v["lid"], "onbekend",
                                        f"Vergunningplicht {v['lid']} niet vast te stellen", ontbr))

    # stap 3: regimes
    regimes = []
    for rg in t.regimes:
        c = ctx()
        r = evalueer(rg["voorwaarde"], c)
        if r is True:
            regimes.append(rg["code"])
        elif is_onbekend(r):
            ontbr = sorted(c.gebruikt_onbekend)
            vragen |= set(ontbr)
            regels.append(Regeluitkomst("regime", rg["code"], "onbekend",
                                        f"Niet vast te stellen of {rg['artikel']} geldt", ontbr))

    # stap 4: zorgplicht
    for z in t.zorgplicht:
        if z["regime"] not in regimes or not (z["zones"] & zones):
            continue
        c = ctx()
        vta = evalueer(z["van_toepassing"], c)
        if vta is False:
            continue
        if is_onbekend(vta):
            ontbr = sorted(c.gebruikt_onbekend)
            vragen |= set(ontbr)
            regels.append(Regeluitkomst("zorgplicht", z["artikel"], "onbekend", z["wsv"], ontbr))
            continue
        if z["soort"] == "aandacht":
            regels.append(Regeluitkomst("zorgplicht", z["artikel"], "aandacht", z["wsv"]))
        elif z["soort"] == "toets" and z["toets"]:
            r = evalueer(z["toets"], c)
            if r is True:
                regels.append(Regeluitkomst("zorgplicht", z["artikel"], "voldoet", z["voorschrift"]))
            elif r is False:
                regels.append(Regeluitkomst("zorgplicht", z["artikel"], "voldoet_niet", z["niet"]))
            else:
                ontbr = sorted(c.gebruikt_onbekend)
                vragen |= set(ontbr)
                regels.append(Regeluitkomst("zorgplicht", z["artikel"], "onbekend", z["wsv"], ontbr))
        else:
            regels.append(Regeluitkomst("zorgplicht", z["artikel"], "voorschrift", z["voorschrift"]))

    soorten = {r.uitkomst for r in regels}
    if "vergunning" in soorten:
        oordeel = "vergunning"
    elif "voldoet_niet" in soorten:
        oordeel = "maatwerk"
    elif "onbekend" in soorten:
        oordeel = "onvolledig"
    elif "aandacht" in soorten:
        oordeel = "aandacht"
    elif soorten & {"voldoet", "voorschrift"}:
        oordeel = "voorschriften"
    else:
        oordeel = "geen_specifieke_regels"
    return Deeluitkomst(oordeel, sorted(zones), regimes, regels, sorted(vragen))


def zwaarste(oordelen) -> str:
    oordelen = list(oordelen)
    return min(oordelen, key=OORDELEN.index) if oordelen else "buiten_beperkingengebied"


# ---------------------------------------------------------------- geometrie
DIJKTYPE = {"waterkerend": "WK", "half-verholen": "HV", "half verholen": "HV", "verholen": "VH"}
ZONETYPE = {"kernzone": "KZ", "beschermingszone": "BZ", "buitenbeschermingszone": "BBZ"}


def zonecode(dijktype, zone) -> str | None:
    a = DIJKTYPE.get(str(dijktype).strip().lower())
    b = ZONETYPE.get(str(zone).strip().lower())
    return f"{a}-{b}" if a and b else None


@dataclass
class GeoLagen:
    zones: Any           # GeoDataFrame met kolommen code, zonecode
    referentielijnen: Any  # GeoDataFrame met kolommen code, dijktype
    kruin: Any           # GeoDataFrame met kolom keringtype (WK/HV/VH)


def laad_geolagen(data_agv: Path) -> GeoLagen:
    import geopandas as gpd
    import pandas as pd
    lv = data_agv / "Legger_verwerkt" / "Alle_keringen"
    zones = gpd.read_file(lv / "Zones_alle_keringen.shp")
    zones["zonecode"] = [zonecode(d, z) for d, z in zip(zones["dijktype"], zones["zone"])]
    zones = zones[zones["zonecode"].notna()].to_crs(28992)
    refl = gpd.read_file(lv / "Referentielijnen_alle_keringen.shp").to_crs(28992)
    kr = []
    for namen, kt in [(("waterkerend_kruin.shp", "19_kruin_waterkerend.shp"), "WK"),
                      (("half-verholen_kruin.shp", "20_kruin_half_verholen.shp"), "HV"),
                      (("verholen_kruin.shp", "21_kruin_verholen.shp"), "VH")]:
        opties = [lv / namen[0], data_agv / "kruinlagen" / namen[1]]
        p = next((o for o in opties if o.exists()), None)
        if p is not None:
            g = gpd.read_file(p).to_crs(28992)
            g = g.explode(index_parts=False)
            g["keringtype"] = kt
            kr.append(g[["keringtype", "geometry"]])
    kruin = gpd.GeoDataFrame(pd.concat(kr, ignore_index=True), crs=28992) if kr else None
    return GeoLagen(zones, refl, kruin)


def _hoek(geom) -> float | None:
    cs = list(geom.coords) if geom.geom_type == "LineString" else None
    if not cs or len(cs) < 2:
        return None
    (x0, y0), (x1, y1) = cs[0][:2], cs[-1][:2]
    if x0 == x1 and y0 == y1:
        return None
    return math.degrees(math.atan2(y1 - y0, x1 - x0)) % 180


def _lokale_richting(lijn, punt) -> tuple[float, tuple[float, float]] | None:
    """Richting (graden) en voetpunt van het stuk referentielijn het dichtst bij punt."""
    from shapely.geometry import LineString
    d = lijn.project(punt)
    a = lijn.interpolate(max(d - 2.0, 0))
    b = lijn.interpolate(min(d + 2.0, lijn.length))
    if a.equals(b):
        return None
    return math.degrees(math.atan2(b.y - a.y, b.x - a.x)) % 180, (lijn.interpolate(d).x, lijn.interpolate(d).y)


def _offsetbereik(geom, richting, voet) -> tuple[float, float]:
    """Bereik van afstanden haaks op de kering (min, max) van alle punten van geom."""
    r = math.radians(richting)
    nx, ny = -math.sin(r), math.cos(r)
    pts = []
    if hasattr(geom, "geoms"):
        for g in geom.geoms:
            pts += list(g.exterior.coords) if g.geom_type == "Polygon" else list(g.coords)
    elif geom.geom_type == "Polygon":
        pts = list(geom.exterior.coords)
    else:
        pts = list(geom.coords)
    offs = [(p[0] - voet[0]) * nx + (p[1] - voet[1]) * ny for p in pts]
    return (min(offs), max(offs)) if offs else (0.0, 0.0)


def _aandeel(deel, vlak, richting, voet, straal=50.0):
    """Welk deel van de breedte van 'vlak' (lokaal rond voet) het deel overbrugt."""
    from shapely.geometry import Point
    lokaal = vlak.intersection(Point(voet).buffer(straal))
    if lokaal.is_empty:
        return None
    vmin, vmax = _offsetbereik(lokaal, richting, voet)
    stuk = deel.intersection(lokaal)
    if stuk.is_empty or vmax - vmin <= 0:
        return 0.0
    dmin, dmax = _offsetbereik(stuk, richting, voet) if stuk.geom_type in ("LineString", "MultiLineString") else (0, 0)
    return max(0.0, min(1.0, (min(dmax, vmax) - max(dmin, vmin)) / (vmax - vmin)))


# ---------------------------------------------------------------- opsplitsen in sub-tracés
# Regels (oefenen_traces/kering_scenarios/regels_opsplitsen.md, afgesproken met Nico):
#  - knip alleen op materiaal/diameter-wissel, aftakkingen en begin/eind van een boring of mantelbuis;
#    een zonegrens of een bocht is geen knip;
#  - elk sub-tracé krijgt de zwaarste zone die het raakt (KZ > BZ > BBZ) en haaks/evenwijdig als kenmerk;
#  - delen buiten alle zones vallen af;
#  - stukjes < KORT_STUK_M zonder andere zone gaan op in de tak waar ze aan vastzitten, behalve naast
#    een boring of mantelbuis;
#  - elke boring is een eigen sub-tracé; alle mantelbuizen samen vormen één sub-tracé;
#  - huisaansluitingen horen bij het sub-tracé waar ze aan vastzitten.
ZONE_ZWAARTE = {"KZ": 0, "BZ": 1, "BBZ": 2}
BOORTERMEN = [  # (patroon, waarde voor invoerveld boormethode; None = methode niet af te leiden)
    (r"gestuurde\s*boring|\bhdd\b", "gestuurde_boring"),
    (r"boogzinker", "boogzinker"),
    (r"boogboring", "boogboring"),
    (r"nano[\s_-]?drill", "nano_drill"),
    (r"raketboring", "raketboring"),
    (r"\bpersing\b", None),
    (r"\bboring\b", None),
]
_OVERLAP_M = 0.1      # een boring/mantelbuis die over een leiding ligt: zoveel marge


def zone_niveau(zonecode: str) -> str:
    return str(zonecode).split("-")[-1]


def zwaarste_zones(zones) -> set[str]:
    """Alleen de zones van het zwaarste niveau (KZ boven BZ boven BBZ)."""
    zones = [z for z in zones if zone_niveau(z) in ZONE_ZWAARTE]
    if not zones:
        return set()
    best = min(ZONE_ZWAARTE[zone_niveau(z)] for z in zones)
    return {z for z in zones if ZONE_ZWAARTE[zone_niveau(z)] == best}


def _teksten(attrs: dict) -> str:
    return " ".join(str(v) for k, v in attrs.items()
                    if not str(k).startswith("__") and str(k).lower() not in ("geometry", "trace", "nr")
                    and isinstance(v, str)).lower()


def _getal(v):
    try:
        f = float(str(v).replace(",", "."))
        return f if f == f else None
    except (TypeError, ValueError):
        return None


def aanlegwijze(attrs: dict) -> tuple[str, str | None]:
    """('open' | 'boring' | 'mantelbuis', boormethode) uit de attributen van een lijn."""
    a = {str(k).lower(): v for k, v in attrs.items()}
    tekst = _teksten(attrs)
    if (not _leeg(a.get("type")) and str(a.get("type")).strip().lower() == "mantelbuis") or "mantel" in tekst:
        return "mantelbuis", None
    bm = "" if _leeg(a.get("boormethode")) else str(a.get("boormethode")).strip().lower()
    if bm and bm != "geen":
        return "boring", bm
    for patroon, waarde in BOORTERMEN:
        if re.search(patroon, tekst):
            return "boring", waarde
    return "open", None


def _leeg(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or str(v).strip().lower() in ("", "nan", "none")


def _eigenschappen(attrs: dict) -> tuple:
    a = {str(k).lower(): v for k, v in attrs.items()}
    mat = "" if _leeg(a.get("materiaal")) else str(a.get("materiaal")).strip().upper()
    mat = mat.split()[0] if mat else ""
    d = _getal(a.get("diameter_mm", a.get("diameter")))
    return mat, d


def _is_ha_attr(attrs: dict) -> bool:
    a = {str(k).lower(): v for k, v in attrs.items()}
    if str(a.get("doel") or "").strip().lower() == "huisaansluiting":
        return True
    return "aansluit" in _teksten(attrs)


@dataclass
class _Lijn:
    id: Any
    groep: Any
    geom: Any            # LineString
    attrs: dict
    aanleg: str
    boormethode: str | None
    mat: str
    diam: float | None
    ha: bool = False

    @property
    def sleutel(self):
        if self.aanleg == "boring":
            return ("boring", self.id)          # elke boring apart
        if self.aanleg == "mantelbuis":
            return ("mantelbuis",)
        return ("open", self.mat, self.diam)


def _delen(geom):
    if geom is None or geom.is_empty:
        return []
    if geom.geom_type == "LineString":
        return [geom]
    if hasattr(geom, "geoms"):
        uit = []
        for g in geom.geoms:
            uit += _delen(g)
        return uit
    return []


def _samen(lijnen):
    from shapely.geometry import MultiLineString
    from shapely.ops import linemerge
    lijnen = [l for l in lijnen if l is not None and not l.is_empty]
    if not lijnen:
        return None
    if len(lijnen) == 1:
        return lijnen[0]
    m = linemerge(MultiLineString([c for l in lijnen for c in _delen(l)]))
    return m


def _knip_op_afstanden(lijn, afstanden, tol):
    from shapely.ops import substring
    ds = sorted({0.0, lijn.length, *[d for d in afstanden if tol < d < lijn.length - tol]})
    return [substring(lijn, a, b) for a, b in zip(ds, ds[1:]) if b - a > 1e-6]


def _zones_geraakt(geom, zones, eps=0.01) -> tuple[set, set]:
    """Zones (en keringcodes) waar geom met meer dan eps meter doorheen loopt."""
    zs, ks = set(), set()
    if geom is None or geom.is_empty or len(zones) == 0:
        return zs, ks
    for code, kering, poly in zip(zones["zonecode"], zones["code"], zones.geometry):
        if geom.intersects(poly) and geom.intersection(poly).length > eps:
            zs.add(code)
            ks.add(kering)
    return zs, ks


def wijze_van_kruisen(hoek, params) -> str | None:
    if hoek is None:
        return None
    return "haaks" if hoek >= params.get("HOEK_MIN", 75) else "evenwijdig"


def splits_traces(lijnen: list[dict], lagen: GeoLagen, params: dict) -> list[dict]:
    """Deelt een tekening (alle lijnen samen) op in sub-tracés volgens de regels hierboven.

    lijnen: dicts met 'id', 'geometry' (LineString/MultiLineString), 'attrs' (dict) en 'groep'
    (bijv. de shapefile: aftakkingen worden alleen binnen een groep gezocht).
    Geeft per sub-tracé: naam, groep, soort (open/boring/mantelbuis/huisaansluiting), geometry (inclusief
    huisaansluitingen), hoofdlijn (zonder), zones (alle geraakte), toetszones (alleen het zwaarste niveau),
    zwaarste_zone (KZ/BZ/BBZ), wijze (haaks/evenwijdig), keringen, bronnen, hoofdbron, materiaal,
    diameter, invoer (wat uit de opsplitsing volgt, zoals boormethode) en geo (geo-feiten).
    """
    from shapely.geometry import Point
    from shapely.ops import unary_union

    tol = params.get("SNAP_AFSTAND_M", 0.5)
    kort = params.get("KORT_STUK_M", 5.0)

    # 1. lijnen inlezen en indelen
    alle: list[_Lijn] = []
    for l in lijnen:
        attrs = dict(l.get("attrs") or {})
        aanleg, bm = aanlegwijze(attrs)
        mat, d = _eigenschappen(attrs)
        for g in _delen(l["geometry"]):
            if g.length > 0:
                alle.append(_Lijn(l["id"], l.get("groep", ""), g, attrs, aanleg, bm, mat, d))
    if not alle:
        return []
    totaal = unary_union([x.geom for x in alle])
    zones = lagen.zones[lagen.zones.intersects(totaal.buffer(1))]
    if len(zones) == 0:
        return []
    # kleine kieren tussen zonevlakken dichten, anders valt een tracé daar onterecht in stukken
    zone_vlak = unary_union(list(zones.geometry)).buffer(tol).buffer(-tol)

    # huisaansluiting: label/doel zegt het, of kleine diameter (<= 63 mm) en kort (<= 30 m);
    # en altijd: één uiteinde zit vast aan een andere (niet-huisaansluiting) lijn van dezelfde groep
    def vast_aan(x, kandidaten):
        for p in (Point(x.geom.coords[0]), Point(x.geom.coords[-1])):
            for y in kandidaten:
                if y is not x and y.groep == x.groep and y.geom.distance(p) <= tol:
                    return p
        return None

    for x in alle:
        if x.aanleg != "open":
            continue
        klein = x.diam is not None and x.diam <= 63 and x.geom.length <= 30
        if _is_ha_attr(x.attrs) or klein:
            andere = [y for y in alle if not (_is_ha_attr(y.attrs) or (y.diam is not None and y.diam <= 63 and y.geom.length <= 30))]
            if vast_aan(x, andere) is not None:
                x.ha = True
    hoofd = [x for x in alle if not x.ha]
    speciaal = [x for x in hoofd if x.aanleg in ("boring", "mantelbuis")]

    # 2. waar een boring of mantelbuis over een open leiding ligt, telt dat stuk als boring/mantelbuis
    stukken: list[_Lijn] = []
    for x in hoofd:
        g = x.geom
        if x.aanleg == "open":
            for s in speciaal:
                if g.intersects(s.geom.buffer(_OVERLAP_M)) and g.intersection(s.geom.buffer(_OVERLAP_M)).length > 2 * _OVERLAP_M + tol:
                    g = g.difference(s.geom.buffer(_OVERLAP_M, cap_style="flat"))
        for d in _delen(g):
            if d.length > 1e-6:
                stukken.append(_Lijn(x.id, x.groep, d, x.attrs, x.aanleg, x.boormethode, x.mat, x.diam))

    # 3. knippen waar een andere lijn (of boring/mantelbuis) met een uiteinde op de lijn aansluit
    eindpunten = [(y, Point(c)) for y in stukken for c in (y.geom.coords[0], y.geom.coords[-1])]
    randen: list[_Lijn] = []
    for x in stukken:
        afst = [x.geom.project(p) for y, p in eindpunten
                if y is not x and (y.groep == x.groep or y.aanleg != "open" or x.aanleg != "open")
                and x.geom.distance(p) <= tol + _OVERLAP_M]
        for d in _knip_op_afstanden(x.geom, afst, tol):
            randen.append(_Lijn(x.id, x.groep, d, x.attrs, x.aanleg, x.boormethode, x.mat, x.diam))

    # 4. knooppunten: uiteinden binnen tol (+ marge rond boring/mantelbuis) zijn één knoop
    #    (per groep: een lijn uit een andere shapefile die toevallig daar eindigt, is geen aftakking)
    knopen: list[tuple] = []
    def knoop(groep, p):
        for i, (g, k) in enumerate(knopen):
            if g == groep and k.distance(p) <= tol + 2 * _OVERLAP_M:
                return i
        knopen.append((groep, p))
        return len(knopen) - 1
    eind = [(knoop(r.groep, Point(r.geom.coords[0])), knoop(r.groep, Point(r.geom.coords[-1]))) for r in randen]
    aan_knoop: dict[int, list[int]] = {}
    for i, (a, b) in enumerate(eind):
        aan_knoop.setdefault(a, []).append(i)
        if b != a:
            aan_knoop.setdefault(b, []).append(i)

    # 5. randen samenvoegen tot takken: alleen door knopen met precies twee randen van dezelfde
    #    groep en met dezelfde eigenschappen (materiaal, diameter, aanlegwijze)
    ouder = list(range(len(randen)))
    def wortel(i):
        while ouder[i] != i:
            ouder[i] = ouder[ouder[i]]
            i = ouder[i]
        return i
    for k, rs in aan_knoop.items():
        if len(rs) == 2:
            a, b = rs
            if randen[a].groep == randen[b].groep and randen[a].sleutel == randen[b].sleutel:
                ouder[wortel(a)] = wortel(b)
    takken: dict[int, list[int]] = {}
    for i in range(len(randen)):
        takken.setdefault(wortel(i), []).append(i)

    # 6. per tak alleen het deel binnen de zones; elk aaneengesloten stuk is een kandidaat
    kand = []
    for leden in takken.values():
        g = _samen([randen[i].geom for i in leden])
        binnen = g.intersection(zone_vlak)
        langste = max(leden, key=lambda i: randen[i].geom.length)
        r0 = randen[langste]
        delen = [d for d in _delen(_samen(_delen(binnen)) if not binnen.is_empty else None) if d.length > 1e-3]
        # stukken die op minder dan tol van elkaar liggen (kleine kier tussen twee lijnen) horen bij elkaar
        groep_van = list(range(len(delen)))
        for i in range(len(delen)):
            for j in range(i):
                if delen[i].distance(delen[j]) <= tol + 2 * _OVERLAP_M:
                    oud, nieuw = groep_van[i], groep_van[j]
                    groep_van = [nieuw if g == oud else g for g in groep_van]
        for gid in sorted(set(groep_van)):
            stuk = [d for d, g in zip(delen, groep_van) if g == gid]
            zs, ks = _zones_geraakt(_samen(stuk), zones)
            if not zs:
                continue
            kand.append(dict(lijnen=stuk, has=[], aanleg=r0.aanleg, boormethode=r0.boormethode,
                             groep=r0.groep, mat=r0.mat, diam=r0.diam,
                             bronnen={randen[i].id for i in leden}, hoofdbron=r0.id,
                             lengte_hoofd=r0.geom.length, zones=zs, keringen=ks))

    # 7. kleine stukjes (< kort) zonder andere zone opnemen in de tak waar ze aan vastzitten,
    #    behalve als ze naast een boring of mantelbuis liggen (die worden altijd apart beoordeeld)
    def geom_van(c):
        return _samen(c["lijnen"] + c["has"])
    def raakt(c1, c2):
        return geom_van(c1).distance(geom_van(c2)) <= tol + 2 * _OVERLAP_M
    speciaal_geom = unary_union([s.geom for s in speciaal]) if speciaal else None
    veranderd = True
    while veranderd:
        veranderd = False
        for c in sorted(kand, key=lambda c: geom_van(c).length):
            g = geom_van(c)
            if c["aanleg"] != "open" or g.length >= kort or c["has"]:
                continue
            if speciaal_geom is not None and g.distance(speciaal_geom) <= tol + 2 * _OVERLAP_M:
                continue
            buren = [b for b in kand if b is not c and b["aanleg"] == "open" and b["groep"] == c["groep"]
                     and c["zones"] <= b["zones"] and raakt(c, b)]
            if not buren:
                continue
            b = max(buren, key=lambda b: geom_van(b).length)
            b["lijnen"] += c["lijnen"]
            b["bronnen"] |= c["bronnen"]
            kand.remove(c)
            veranderd = True
            break

    # 8. alle mantelbuizen samen één sub-tracé
    mb = [c for c in kand if c["aanleg"] == "mantelbuis"]
    if len(mb) > 1:
        eerste = mb[0]
        for c in mb[1:]:
            eerste["lijnen"] += c["lijnen"]
            eerste["bronnen"] |= c["bronnen"]
            eerste["zones"] |= c["zones"]
            eerste["keringen"] |= c["keringen"]
            kand.remove(c)
        eerste["groep"] = "mantelbuizen"

    # 9. huisaansluitingen bij het sub-tracé waar ze aan vastzitten
    for x in [x for x in alle if x.ha]:
        binnen = x.geom.intersection(zone_vlak)
        if binnen.is_empty or binnen.length <= 1e-3:
            continue
        p = vast_aan(x, hoofd)
        doel = None
        if p is not None:
            dichtst = [c for c in kand if geom_van(c).distance(p) <= tol + 2 * _OVERLAP_M]
            if dichtst:
                doel = min(dichtst, key=lambda c: geom_van(c).distance(p))
        if doel is None:   # hoofdleiding ligt buiten de zones: eigen sub-tracé
            zs, ks = _zones_geraakt(binnen, zones)
            if not zs:
                continue
            kand.append(dict(lijnen=[], has=[binnen], aanleg="huisaansluiting", boormethode=None, groep=x.groep,
                             mat=x.mat, diam=x.diam, bronnen={x.id}, hoofdbron=x.id, lengte_hoofd=x.geom.length,
                             zones=zs, keringen=ks))
            continue
        doel["has"].append(binnen)
        doel["bronnen"].add(x.id)

    # 10. kenmerken en geo-feiten per sub-tracé
    uit = []
    volgorde = {"open": 0, "boring": 1, "huisaansluiting": 2, "mantelbuis": 3}
    kand.sort(key=lambda c: (volgorde[c["aanleg"]], str(c["groep"]),
                             -geom_van(c).bounds[1] if c["lijnen"] else 0, geom_van(c).bounds[0] if c["lijnen"] else 0))
    for c in kand:
        hoofdlijn = _samen(c["lijnen"]) if c["lijnen"] else _samen(c["has"])
        alles = _samen(c["lijnen"] + c["has"])
        zs, ks = _zones_geraakt(alles, zones)
        zs |= c["zones"]
        ks |= c["keringen"]
        toets = zwaarste_zones(zs)
        geo = geo_feiten(hoofdlijn, zs, ks, lagen, params)
        if geo.get("hoek_met_kering") is None and hoofdlijn.geom_type != "LineString":
            # meerdere losse stukken (bijv. mantelbuizen): richting van het langste stuk
            langste = max(_delen(hoofdlijn), key=lambda g: g.length)
            geo["hoek_met_kering"] = geo_feiten(langste, zs, ks, lagen, params).get("hoek_met_kering")
        invoer = {}
        if c["aanleg"] == "mantelbuis":
            invoer["type"] = "mantelbuis"
        elif c["aanleg"] == "boring":
            if c["boormethode"]:
                invoer["boormethode"] = c["boormethode"]
        elif c["aanleg"] == "open":
            invoer["boormethode"] = "geen"
        elif c["aanleg"] == "huisaansluiting":
            invoer["doel"] = "huisaansluiting"
        uit.append(dict(groep=c["groep"], soort=c["aanleg"], geometry=alles, hoofdlijn=hoofdlijn,
                        zones=zs, toetszones=toets,
                        zwaarste_zone=zone_niveau(next(iter(toets))) if toets else None,
                        wijze=wijze_van_kruisen(geo.get("hoek_met_kering"), params),
                        keringen=sorted(ks), bronnen=sorted(c["bronnen"], key=str), hoofdbron=c["hoofdbron"],
                        materiaal=c["mat"] or None, diameter=c["diam"], huisaansluitingen=len(c["has"]),
                        invoer=invoer, geo=geo))
    for i, d in enumerate(uit, 1):
        d["nr"] = i
        d["naam"] = f"{d['groep']}#{i}"
    return uit


def splits_trace(lijn, lagen: GeoLagen, params: dict, attrs: dict | None = None) -> list[dict]:
    """Eén losse lijn opsplitsen (oude aanroep). Voor een hele tekening: splits_traces."""
    return splits_traces([dict(id=0, geometry=lijn, attrs=attrs or {}, groep="")], lagen, params)


def geo_feiten(deel, zones, keringen, lagen: GeoLagen, params: dict) -> dict:
    from shapely.ops import unary_union
    zonder_kruin = lagen.kruin is None
    feiten = dict(kruist_referentielijn=False, dwars_aandeel_kernzone=0.0, hoek_met_kering=None,
                  raakt_kruin=None if zonder_kruin else False, dwars_door_kruin=None if zonder_kruin else False)
    if not zones:
        return feiten
    refl = lagen.referentielijnen
    refl = refl[refl["code"].isin(keringen)] if keringen else refl.iloc[0:0]
    if len(refl) == 0:
        refl = lagen.referentielijnen[lagen.referentielijnen.distance(deel) < 100]
    if len(refl):
        lijn = unary_union(list(refl.geometry))
        feiten["kruist_referentielijn"] = bool(deel.intersects(lijn))
        from shapely.ops import linemerge
        if lijn.geom_type == "MultiLineString":
            lijn = linemerge(lijn)
        if lijn.geom_type == "MultiLineString":
            lijn = min(lijn.geoms, key=lambda g: g.distance(deel))
        mid = deel.interpolate(0.5, normalized=True)
        lr = _lokale_richting(lijn, mid)
        hd = _hoek(deel) if deel.geom_type == "LineString" else None
        if lr and hd is not None:
            verschil = abs(hd - lr[0]) % 180
            feiten["hoek_met_kering"] = round(min(verschil, 180 - verschil), 1)
        if lr:
            richting, _ = lr
            voet_d = lijn.project(mid)
            vp = lijn.interpolate(voet_d)
            voet = (vp.x, vp.y)
            kz = lagen.zones[(lagen.zones["zonecode"] == "WK-KZ") & (lagen.zones["code"].isin(keringen))]
            if len(kz) and "WK-KZ" in zones:
                kzv = unary_union(list(kz.geometry))
                # meet bij het midden en bij elk punt waar het deel de referentielijn kruist; neem het grootste
                punten = [voet]
                kruis = deel.intersection(lijn)
                for p in getattr(kruis, "geoms", [kruis]):
                    if p.geom_type == "Point":
                        punten.append((p.x, p.y))
                beste = None
                for pt in punten:
                    from shapely.geometry import Point as _P
                    lr2 = _lokale_richting(lijn, _P(pt))
                    if not lr2:
                        continue
                    a = _aandeel(deel, kzv, lr2[0], lr2[1])
                    if a is not None and (beste is None or a > beste):
                        beste = a
                feiten["dwars_aandeel_kernzone"] = round(beste, 2) if beste is not None else None
            if lagen.kruin is not None:
                kt = {z.split("-")[0] for z in zones}
                kr = lagen.kruin[lagen.kruin["keringtype"].isin(kt) & lagen.kruin.intersects(deel.buffer(0.01))]
                if len(kr):
                    kv = unary_union(list(kr.geometry))
                    feiten["raakt_kruin"] = True
                    a = _aandeel(deel, kv, richting, voet)
                    feiten["dwars_door_kruin"] = bool(a is not None and a >= params.get("DWARS_AANDEEL", 0.75))
    return feiten


def toets_subtrace(t: Template, d: dict, invoer: dict) -> Deeluitkomst:
    """Toetst één sub-tracé op zijn zwaarste zone. Wat uit de opsplitsing volgt (boring, mantelbuis) gaat voor."""
    return toets_deel(t, {**invoer, **d["invoer"]}, d["toetszones"], d["geo"])


def toets_traces(t: Template, lijnen: list[dict], lagen: GeoLagen, invoer_van=lambda d: {}):
    """Splitst een tekening op en toetst elk sub-tracé. invoer_van(sub-tracé) geeft de invoer."""
    delen = splits_traces(lijnen, lagen, t.params)
    resultaten = [(d, toets_subtrace(t, d, invoer_van(d))) for d in delen]
    return zwaarste(r.oordeel for _, r in resultaten), resultaten


def toets_trace(t: Template, lijn, invoer: dict, lagen: GeoLagen):
    return toets_traces(t, [dict(id=0, geometry=lijn, attrs={}, groep="")], lagen, lambda d: invoer)


# ---------------------------------------------------------------- testgevallen
def _lees_kv(s) -> dict:
    uit = {}
    for paar in str(s or "").split(";"):
        if "=" not in paar:
            continue
        k, v = (x.strip() for x in paar.split("=", 1))
        if v in ("True", "False"):
            v = v == "True"
        else:
            try:
                v = float(v) if "." in v else int(v)
            except ValueError:
                pass
        uit[k] = v
    return uit


def draai_testgevallen(pad, schrijf=True) -> int:
    t = laad_template(pad)
    wb = load_workbook(pad)
    ws = wb["Testgevallen"]
    koppen = [c.value for c in ws[1]]
    k = {n: i for i, n in enumerate(koppen)}
    fout = 0
    for row in ws.iter_rows(min_row=2):
        if not row[0].value:
            continue
        zones = {z.strip() for z in str(row[k["zone"]].value or "").split(",") if z.strip()}
        r = toets_deel(t, _lees_kv(row[k["eigenschappen"]].value), zones, _lees_kv(row[k["geo-feiten"]].value))
        verwacht = str(row[k["verwacht oordeel"]].value).strip()
        ok = r.oordeel == verwacht
        fout += not ok
        arts = []
        for soort in ("vergunning", "voldoet_niet", "onbekend", "aandacht"):
            a = r.artikelen(soort)
            if a:
                arts.append(f"{soort}: {', '.join(a)}")
        if r.vragen:
            arts.append("vragen: " + ", ".join(r.vragen))
        row[k["uitkomst rekenkern"]].value = r.oordeel
        row[k["geraakte artikelen"]].value = " | ".join(arts)
        row[k["status"]].value = "OK" if ok else "AFWIJKING"
        print(f"{row[0].value}: {'OK ' if ok else 'XX '} verwacht={verwacht:<24} uitkomst={r.oordeel:<24} {' | '.join(arts)}")
    if schrijf:
        wb.save(pad)
    print(f"{fout} afwijking(en)")
    return fout


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[2] == "--test":
        sys.exit(1 if draai_testgevallen(sys.argv[1]) else 0)
    print(__doc__)
