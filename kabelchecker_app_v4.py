"""Kabelchecker v4 - toetst tracés aan de Waterschapsverordening AGV met regeltemplate v4.

Starten:  streamlit run kabelchecker_app_v4.py
"""
from __future__ import annotations

import io
import re
import tempfile
import zipfile
from pathlib import Path

import geopandas as gpd
import pandas as pd
import streamlit as st

import kabelchecker_v4_engine as K

APP_VERSION = "4.2"
HIER = Path(__file__).resolve().parent
STANDAARD_TEMPLATE = HIER / "regeltemplate_v4_AGV.xlsx"
_KANDIDATEN = [HIER.parent / "Kaartlezer" / "data_AGV", HIER.parent.parent / "Kaartlezer" / "data_AGV",
               HIER.parent / "data_AGV", HIER / "data_AGV"]
STANDAARD_DATA_AGV = next((p for p in _KANDIDATEN if p.exists()), _KANDIDATEN[0])

OORDEEL_TEKST = {
    "vergunning": "Vergunningplichtig",
    "maatwerk": "Maatwerk nodig",
    "onvolledig": "Onvolledig: gegevens ontbreken",
    "aandacht": "Aandachtspunt voor vergunningverlener",
    "voorschriften": "Toegestaan met voorschriften",
    "geen_specifieke_regels": "Geen specifieke regels van toepassing",
    "buiten_beperkingengebied": "Buiten het beperkingengebied",
}
OORDEEL_KLEUR = {
    "vergunning": [200, 30, 30, 255],
    "maatwerk": [230, 120, 0, 255],
    "onvolledig": [150, 90, 200, 255],
    "aandacht": [220, 180, 0, 255],
    "voorschriften": [40, 150, 60, 255],
    "geen_specifieke_regels": [90, 90, 90, 255],
    "buiten_beperkingengebied": [120, 120, 120, 255],
}


# ------------------------------------------------------------------ inlezen
@st.cache_resource(show_spinner="Leggerlagen AGV laden...")
def geolagen(pad: str):
    return K.laad_geolagen(Path(pad))


def lees_traces(uploads) -> gpd.GeoDataFrame:
    frames = []
    shp_bundel: dict[str, bytes] = {}
    for up in uploads:
        naam, data = up.name, up.getvalue()
        ext = Path(naam).suffix.lower()
        if ext in (".geojson", ".json"):
            with tempfile.NamedTemporaryFile(suffix=".geojson", delete=False) as f:
                f.write(data)
            g = gpd.read_file(f.name)
            g["__bestand__"] = Path(naam).stem
            frames.append(g)
        elif ext in (".shp", ".dbf", ".shx", ".prj", ".cpg"):
            shp_bundel[naam] = data
        elif ext == ".zip":
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for m in zf.namelist():
                    if Path(m).suffix.lower() in (".shp", ".dbf", ".shx", ".prj", ".cpg"):
                        shp_bundel[Path(m).name] = zf.read(m)
    for shp in [n for n in shp_bundel if n.lower().endswith(".shp")]:
        stam = Path(shp).stem
        with tempfile.TemporaryDirectory() as td:
            for n, raw in shp_bundel.items():
                if Path(n).stem == stam:
                    (Path(td) / n).write_bytes(raw)
            g = gpd.read_file(Path(td) / shp)
        g["__bestand__"] = stam
        frames.append(g)
    if not frames:
        return gpd.GeoDataFrame(geometry=[], crs=28992)
    uit = []
    for g in frames:
        if g.crs is None or (g.total_bounds[0] > 1000 and g.crs.to_epsg() != 28992 and g.crs.is_geographic):
            g = g.set_crs(28992, allow_override=True)
        uit.append(g.to_crs(28992))
    gdf = gpd.GeoDataFrame(pd.concat(uit, ignore_index=True), crs=28992)
    gdf = gdf[gdf.geometry.notna() & gdf.geom_type.isin(["LineString", "MultiLineString"])].reset_index(drop=True)
    gdf["trace"] = [f"{b}#{i + 1}" for i, b in enumerate(gdf["__bestand__"])]
    return gdf


# ------------------------------------------------------------------ attributen uit Kaartlezer-shapefiles
MEDIUM_UIT_SOORT = {"water": "drinkwater", "dw": "drinkwater", "drinkwater": "drinkwater", "hwa": "hwa", "rwa": "hwa",
                    "dwa": "dwa", "vwa": "dwa", "gwa": "dwa", "pers": "dwa", "gas": "gas", "g": "gas"}


def voorinvullen(rij: pd.Series, velden: list[str]) -> dict:
    """Haalt wat kan uit de attributen van een tracé (Kaartlezer-uitvoer of eigen GeoJSON)."""
    attr = {str(k).lower(): v for k, v in rij.items() if k not in ("geometry",)}
    uit = {v: attr.get(v) for v in velden if attr.get(v) not in (None, "")}
    soort = str(attr.get("soort") or "").strip().lower()
    if "type" not in uit or str(uit.get("type")).lower() not in ("kabel", "leiding", "mantelbuis"):
        uit.pop("type", None)
        if "mantel" in soort:
            uit["type"] = "mantelbuis"
        elif any(k in soort for k in ("kabel", "elec", "ls", "ms", "telecom", "glas")):
            uit["type"] = "kabel"
        elif soort:
            uit["type"] = "leiding"
    if "medium" not in uit and soort:
        for k, m in MEDIUM_UIT_SOORT.items():
            if soort == k or soort.startswith(k):
                uit["medium"] = m
                break
    mat = str(attr.get("materiaal") or "")
    if mat:
        eerste = mat.split()[0].upper()
        uit["materiaal"] = {"PE": "PE", "PVC": "PVC", "PP": "PP", "GN": "staal", "GGIJ": "staal", "STAAL": "staal",
                            "AC": "AC", "BETON": "beton", "GRES": "gres", "KOPER": "koper", "GVK": "GVK"}.get(eerste, eerste)
        m = re.search(r"SDR\s*(\d+)", mat, re.I)
        if m:
            uit["sdr"] = int(m.group(1))
    if "diameter_mm" not in uit and attr.get("diameter") not in (None, ""):
        try:
            uit["diameter_mm"] = float(attr["diameter"])
        except (TypeError, ValueError):
            pass
    status = str(attr.get("status") or "").lower()
    if "werkwijze" not in uit and status:
        if "schuim" in status:
            uit["werkwijze"] = "volschuimen"
        elif "nieuw" in status:
            uit["werkwijze"] = "nieuw"
        elif any(k in status for k in ("vervallen", "verwijder")):
            uit["werkwijze"] = "verwijderen"
    if soort == "pers" and "druk_bar" not in uit:
        pass  # persleiding: druk vraagt de app (afspraak: persleiding > 1 bar), niet aannemen
    return {k: v for k, v in uit.items() if v not in (None, "")}


# ------------------------------------------------------------------ teksten en kleuren voor de resultaten
OORDEEL_HEX = {o: "#{:02x}{:02x}{:02x}".format(*k[:3]) for o, k in OORDEEL_KLEUR.items()}
OORDEEL_STIP = {"vergunning": "🔴", "maatwerk": "🟠", "onvolledig": "🟣", "aandacht": "🟡",
                "voorschriften": "🟢", "geen_specifieke_regels": "⚪", "buiten_beperkingengebied": "⚪"}
UITKOMST_TEKST = {"vergunning": "vergunningplichtig", "voldoet": "voldoet", "voldoet_niet": "voldoet niet",
                  "onbekend": "gegevens ontbreken", "voorschrift": "voorschrift", "aandacht": "aandachtspunt"}
UITKOMST_VOLGORDE = {"vergunning": 0, "voldoet_niet": 1, "onbekend": 2, "aandacht": 3, "voldoet": 4, "voorschrift": 5}
SOORT_TEKST = {"open": "open ontgraving", "boring": "boring", "mantelbuis": "mantelbuis", "huisaansluiting": "huisaansluiting"}
# zonering zoals op de leggerkaart van AGV: kernzone rood gearceerd, beschermingszone blauw gearceerd,
# buitenbeschermingszone groene stippen
ZONE_STIJL = {"KZ": ("Kernzone", "#d7191c"), "BZ": ("Beschermingszone", "#1f5fbf"), "BBZ": ("Buitenbeschermingszone", "#1a9641")}
ALTIJD_VRAGEN = ["doel"]  # huisaansluiting of distributie: altijd door de vergunningverlener laten kiezen
PDOK_BRT = "https://service.pdok.nl/brt/achtergrondkaart/wmts/v2_0/grijs/EPSG:3857/{z}/{x}/{y}.png"
PDOK_LUCHTFOTO = "https://service.pdok.nl/hwh/luchtfotorgb/wmts/v1_0/Actueel_orthoHR/EPSG:3857/{z}/{x}/{y}.jpeg"
PDOK_ATTR = 'Kaart: <a href="https://www.pdok.nl">PDOK</a> · legger AGV'


def nr(naam: str) -> str:
    """Korte tracénaam zoals overal in de app: #3."""
    return "#" + str(naam).split("#")[-1]


def nrs_tekst(namen) -> str:
    n = [nr(x) for x in namen]
    return n[0] if len(n) == 1 else ", ".join(n[:-1]) + " en " + n[-1]


@st.cache_data
def zone_omschrijving(tpl_pad: str) -> dict[str, str]:
    """Zonecode -> omschrijving uit het tabblad Zones van het regeltemplate."""
    import openpyxl
    try:
        ws = openpyxl.load_workbook(tpl_pad, read_only=True)["Zones"]
        uit = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[3] and re.fullmatch(r"[A-Z]+-[A-Z]+", str(r[0]).strip()):
                uit[str(r[0]).strip()] = str(r[3]).strip()
        return uit
    except Exception:
        return {}


def zone_kort(code: str) -> str:
    """WK-KZ -> 'Kernzone (waterkerend)'."""
    dijk = {"WK": "waterkerend", "HV": "half-verholen", "VH": "verholen"}
    a, _, b = str(code).partition("-")
    return f"{ZONE_STIJL[b][0]} ({dijk.get(a, a)})" if b in ZONE_STIJL else str(code)


def keringen_tekst(codes, keringnaam: dict) -> str:
    """Keringcodes met hun naam, codes met dezelfde naam samen: A541-001 en A541-002 (Kramatweg)."""
    per_naam: dict[str, list[str]] = {}
    for k in codes:
        per_naam.setdefault(keringnaam.get(k, ""), []).append(k)
    delen = []
    for naam, ks in per_naam.items():
        c = ks[0] if len(ks) == 1 else ", ".join(ks[:-1]) + " en " + ks[-1]
        delen.append(f"{c} ({naam})" if naam else c)
    return "; ".join(delen)


def badge(oordeel: str) -> str:
    return (f"<span style='background:{OORDEEL_HEX[oordeel]};color:white;padding:2px 10px;border-radius:12px;"
            f"font-size:0.85em;white-space:nowrap'>{OORDEEL_TEKST[oordeel]}</span>")


# ------------------------------------------------------------------ kaart (folium, met AGV-arcering)
def _legenda_html() -> str:
    def vlak(css):
        return f"<span style='display:inline-block;width:22px;height:14px;margin-right:6px;vertical-align:middle;{css}'></span>"
    rijen = []
    for lvl, (naam, kl) in ZONE_STIJL.items():
        if lvl == "BBZ":
            css = f"border:1px solid {kl};background:radial-gradient({kl} 1.2px, transparent 1.6px) 0 0/5px 5px"
        else:
            css = f"border:1px solid {kl};background:repeating-linear-gradient(-45deg,{kl} 0 1.5px,transparent 1.5px 5px)"
        rijen.append(vlak(css) + naam)
    lijnen = [f"<span style='display:inline-block;width:22px;height:0;border-top:4px solid {OORDEEL_HEX[o]};"
              f"margin-right:6px;vertical-align:middle'></span>{OORDEEL_TEKST[o]}"
              for o in ["vergunning", "maatwerk", "onvolledig", "aandacht", "voorschriften", "geen_specifieke_regels"]]
    lijnen.append("<span style='display:inline-block;width:22px;height:0;border-top:3px dashed #888;"
                  "margin-right:6px;vertical-align:middle'></span>Uitgezet (telt niet mee)")
    return ("<div style='position:fixed;bottom:18px;left:12px;z-index:9999;background:rgba(255,255,255,.93);"
            "padding:8px 10px;border-radius:6px;box-shadow:0 1px 4px rgba(0,0,0,.3);font:11px/1.5 sans-serif'>"
            "<b>Zones (legger AGV)</b><br>" + "<br>".join(rijen) +
            "<br><b>Oordeel per tracé</b><br>" + "<br>".join(lijnen) + "</div>")


def kaart(sub_van: dict, uitkomsten: dict, uit: set, lagen, zonenaam: dict, keringnaam: dict):
    import folium
    from folium.plugins.pattern import CirclePattern, StripePattern
    import streamlit.components.v1 as components

    alle = gpd.GeoDataFrame({"naam": list(sub_van)}, geometry=[d["geometry"] for d in sub_van.values()], crs=28992)
    b = alle.total_bounds
    m = folium.Map(tiles=None, control_scale=True, max_zoom=21)
    folium.TileLayer(PDOK_BRT, attr=PDOK_ATTR, name="Kaart (BRT grijs)", max_native_zoom=19, max_zoom=21).add_to(m)
    folium.TileLayer(PDOK_LUCHTFOTO, attr=PDOK_ATTR, name="Luchtfoto", max_native_zoom=20, max_zoom=21,
                     show=False, overlay=False).add_to(m)

    patronen = {
        "KZ": StripePattern(angle=-45, weight=1.5, space_weight=4.5, color=ZONE_STIJL["KZ"][1], opacity=0.8).add_to(m),
        "BZ": StripePattern(angle=-45, weight=1.5, space_weight=4.5, color=ZONE_STIJL["BZ"][1], opacity=0.7).add_to(m),
        "BBZ": CirclePattern(width=7, height=7, radius=1.3, weight=0, color=ZONE_STIJL["BBZ"][1],
                             fill_color=ZONE_STIJL["BBZ"][1], opacity=0.8, fill_opacity=0.8).add_to(m),
    }
    zones = lagen.zones.cx[b[0] - 150:b[2] + 150, b[1] - 150:b[3] + 150]
    for lvl in ["BBZ", "BZ", "KZ"]:  # zwaarste zone bovenop
        z = zones[zones["zonecode"].str.endswith("-" + lvl)]
        if z.empty:
            continue
        z = z.assign(zonenaam=[zonenaam.get(c, c) for c in z["zonecode"]], kering=z["code"],
                     keringnaam=[str(n or "") for n in z.get("naam", [""] * len(z))])
        z = z[["kering", "keringnaam", "zonenaam", "geometry"]].to_crs(4326)
        pat, kl = patronen[lvl], ZONE_STIJL[lvl][1]
        folium.GeoJson(
            z.to_json(), name=ZONE_STIJL[lvl][0],
            style_function=lambda f, pat=pat, kl=kl: {"fillPattern": pat, "color": kl, "weight": 1, "fillOpacity": 1},
            tooltip=folium.GeoJsonTooltip(["kering", "keringnaam", "zonenaam"], aliases=["Kering", "", "Zone"]),
        ).add_to(m)

    lijnen = folium.FeatureGroup(name="Tracés").add_to(m)
    labels = folium.FeatureGroup(name="Tracénummers").add_to(m)
    geplaatst = []
    for naam, (oordeel, delen) in uitkomsten.items():
        d = sub_van[naam]
        aan = naam not in uit
        props = {"tracé": nr(naam), "soort": SOORT_TEKST.get(d["soort"], d["soort"]),
                 "lengte": f"{d['hoofdlijn'].length:.0f} m",
                 "kering": keringen_tekst(d.get("keringen", []), keringnaam) or "-",
                 "zone": ", ".join(zonenaam.get(z, z) for z in sorted(d["toetszones"])) or "-",
                 "oordeel": OORDEEL_TEKST[oordeel] if aan else "uitgezet, telt niet mee"}
        g = gpd.GeoDataFrame([props], geometry=[d["geometry"]], crs=28992).to_crs(4326)
        kl = OORDEEL_HEX[oordeel] if aan else "#888888"
        folium.GeoJson(
            g.to_json(),
            style_function=lambda f, kl=kl, aan=aan: {"color": kl, "weight": 5 if aan else 3,
                                                      "dashArray": None if aan else "6 6", "opacity": 0.95},
            highlight_function=lambda f: {"weight": 8},
            tooltip=folium.GeoJsonTooltip(list(props), aliases=["Tracé", "Soort", "Lengte", "Kering", "Zone", "Oordeel"]),
        ).add_to(lijnen)
        h = d["hoofdlijn"]
        mp = h.interpolate(0.5, normalized=True) if h.geom_type == "LineString" else h.representative_point()
        buren = sum(1 for q in geplaatst if q.distance(mp) < 15)  # labels dicht bij elkaar onder elkaar
        geplaatst.append(mp)
        p = gpd.GeoSeries([mp], crs=28992).to_crs(4326).iloc[0]
        folium.Marker([p.y, p.x], icon=folium.DivIcon(
            icon_size=(40, 20), icon_anchor=(20, 10 - 22 * buren),
            html=f"<div style='font:700 13px sans-serif;background:white;border:2px solid {kl};border-radius:4px;"
                 f"text-align:center;color:#222;{'' if aan else 'opacity:.6;'}'>{nr(naam)}</div>")).add_to(labels)

    m.get_root().html.add_child(folium.Element(_legenda_html()))
    folium.LayerControl(collapsed=True).add_to(m)
    wgs = gpd.GeoSeries.from_xy([b[0], b[2]], [b[1], b[3]], crs=28992).to_crs(4326)
    m.fit_bounds([[wgs.y.min(), wgs.x.min()], [wgs.y.max(), wgs.x.max()]], padding=(30, 30))
    components.html(m.get_root().render(), height=620)


# ------------------------------------------------------------------ uitleg per tracé
def waarom(oordeel, d, r, zonenaam, keringnaam, regime_artikel) -> list[str]:
    """Korte zinnen in gewone taal: waar ligt het tracé en waarom dit oordeel."""
    zones = " en ".join(zonenaam.get(z, z).lower() for z in sorted(d["toetszones"])) or "geen zone"
    kering = keringen_tekst(d.get("keringen", []), keringnaam)
    wijze = {"haaks": "kruist haaks", "evenwijdig": "ligt evenwijdig"}.get(d.get("wijze"), "ligt")
    zinnen = [f"Het tracé {wijze} in de {zones}" + (f" van kering {kering}." if kering else ".")]
    if oordeel == "vergunning":
        zinnen += [f"Vergunningplichtig (art. {x.artikel}): {x.tekst}" for x in r.regels if x.uitkomst == "vergunning"]
    elif oordeel == "maatwerk":
        zinnen += [f"Voldoet niet aan art. {x.artikel}: {x.tekst}" for x in r.regels if x.uitkomst == "voldoet_niet"]
    elif oordeel == "onvolledig":
        verg = [x.artikel for x in r.regels if x.stap == "vergunning" and x.uitkomst == "onbekend"]
        reg = [regime_artikel.get(x.artikel, x.artikel) for x in r.regels if x.stap == "regime" and x.uitkomst == "onbekend"]
        zp = [x.artikel for x in r.regels if x.stap == "zorgplicht" and x.uitkomst == "onbekend"]
        delen = []
        if verg:
            delen.append("of er een vergunning nodig is (art. " + ", ".join(verg) + ")")
        if reg:
            delen.append("welke regels gelden (art. " + ", ".join(dict.fromkeys(reg)) + ")")
        if zp:
            delen.append("of aan art. " + ", ".join(dict.fromkeys(zp)) + " wordt voldaan")
        zinnen.append("Nog niet te beoordelen: er ontbreken gegevens om vast te stellen " + " en ".join(delen) + "."
                      if delen else "Nog niet te beoordelen: er ontbreken gegevens.")
    elif oordeel == "aandacht":
        zinnen += [f"Aandachtspunt (art. {x.artikel}): {x.tekst}" for x in r.regels if x.uitkomst == "aandacht"]
    elif oordeel == "voorschriften":
        arts = ", ".join(dict.fromkeys(a for a in (str(regime_artikel.get(g, g)) for g in r.regimes) if a[:1].isdigit()))
        zinnen.append("Toegestaan zonder vergunning" + (f" onder art. {arts}" if arts else "") +
                      ", mits aan de voorschriften hieronder wordt voldaan.")
    elif oordeel == "geen_specifieke_regels":
        zinnen.append("Voor dit tracé gelden in deze zone geen specifieke regels.")
    return zinnen


def excel_rapport(overzicht: pd.DataFrame, regels: pd.DataFrame, voorschriften: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for naam, df in [("Overzicht", overzicht), ("Voorschriften", voorschriften), ("Alle regels", regels)]:
            df.to_excel(xw, sheet_name=naam, index=False)
            ws = xw.sheets[naam]
            ws.freeze_panes = "A2"
            for i, kol in enumerate(df.columns, 1):
                breed = max([len(str(kol))] + [len(str(v)) for v in df[kol].head(200)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(10, breed + 2), 80)
    return buf.getvalue()


# ------------------------------------------------------------------ app
def wachtwoord_ok() -> bool:
    """Online versie: vraag het wachtwoord uit st.secrets. Lokaal (geen secrets) geen wachtwoord."""
    try:
        juist = st.secrets.get("wachtwoord")
    except Exception:
        juist = None
    if not juist or st.session_state.get("ingelogd"):
        return True
    st.title("Kabelchecker")
    with st.form("login"):
        w = st.text_input("Wachtwoord", type="password")
        if st.form_submit_button("Inloggen", type="primary"):
            if w == juist:
                st.session_state["ingelogd"] = True
                st.rerun()
            st.error("Onjuist wachtwoord.")
    return False


def nieuwe_casus():
    """Wist upload, antwoorden en vinkjes; het inloggen blijft staan."""
    casus = st.session_state.get("casus", 0) + 1
    for k in list(st.session_state):
        if k != "ingelogd":
            del st.session_state[k]
    st.session_state["casus"] = casus


def main():
    st.set_page_config(page_title="Kabelchecker", layout="wide")
    if not wachtwoord_ok():
        st.stop()
    casus = st.session_state.setdefault("casus", 0)

    with st.sidebar:
        st.button("Nieuwe casus", on_click=nieuwe_casus, type="primary", width="stretch",
                  help="Wist de upload, de antwoorden en de vinkjes, zodat je met een nieuwe shapefile kunt beginnen.")
        st.header("Tracé uploaden")
        uploads = st.file_uploader("GeoJSON, shapefile-delen of zip", accept_multiple_files=True, key=f"upload_{casus}",
                                   type=["geojson", "json", "shp", "dbf", "shx", "prj", "cpg", "zip"])
        tpl_pad = STANDAARD_TEMPLATE
        with st.expander("Instellingen voor beheer"):
            data_agv = st.text_input("Map data_AGV", str(STANDAARD_DATA_AGV))
            try:
                t = K.laad_template(tpl_pad)
            except Exception as e:
                st.error(f"Template niet te lezen: {e}")
                st.stop()
            t.params["HOEK_MIN"] = st.slider("Minimale hoek voor 'loodrecht' (graden)", 45, 90,
                                             int(t.params.get("HOEK_MIN", 75)))
            t.params["DWARS_AANDEEL"] = st.slider("Aandeel kernzonebreedte voor 'dwars door'", 0.5, 1.0,
                                                  float(t.params.get("DWARS_AANDEEL", 0.75)), 0.05)
            st.caption(f"Regeltemplate: {tpl_pad.name} ({len(t.zorgplicht)} zorgplichtregels, "
                       f"{len(t.vergunning)} vergunningregels)")

    st.title(f"Kabelchecker {APP_VERSION}")
    st.caption("Toetsing van tracés aan de Waterschapsverordening AGV (art. 2.15 en 3.14-3.18), per deeltracé.")

    if not Path(data_agv).exists():
        st.error(f"Map data_AGV niet gevonden: {data_agv}")
        st.stop()
    if not uploads:
        st.info("Upload links een tracé om te beginnen. De lijnen worden opgesplitst in deeltracés: een knip bij een "
                "wissel van materiaal of diameter, bij een aftakking en aan begin en eind van een boring of mantelbuis. "
                "Delen buiten de zones vallen af.")
        st.stop()

    lagen = geolagen(data_agv)
    lijnen = lees_traces(uploads)
    if lijnen.empty:
        st.warning("Geen lijnen gevonden in de upload.")
        st.stop()
    # opsplitsen in sub-tracés (regels_opsplitsen.md); elk sub-tracé wordt apart getoetst
    rij_van = {r["trace"]: r for _, r in lijnen.iterrows()}
    subtraces = K.splits_traces(
        [dict(id=r["trace"], geometry=r.geometry, groep=r["__bestand__"],
              attrs={k: v for k, v in r.items() if k != "geometry"}) for _, r in lijnen.iterrows()],
        lagen, t.params)
    if not subtraces:
        st.success("Geen enkel tracé raakt een zone van de legger: buiten het beperkingengebied.")
        st.stop()
    sub_van = {d["naam"]: d for d in subtraces}

    zonenaam = zone_omschrijving(str(tpl_pad))
    keringnaam = {}
    if "naam" in lagen.zones.columns:
        keringnaam = {str(c): str(n) for c, n in zip(lagen.zones["code"], lagen.zones["naam"]) if n}
    regime_artikel = {r["code"]: r["artikel"] for r in t.regimes}

    velden = [str(v["veld"]) for v in t.invoervelden]
    keuzes = {str(v["veld"]): [x.strip() for x in str(v.get("waarden") or "").split(";")]
              for v in t.invoervelden if str(v.get("type")) == "keuze"}
    info = {str(v["veld"]): v for v in t.invoervelden}
    basis = {d["naam"]: {**voorinvullen(rij_van[d["hoofdbron"]], velden), **d["invoer"]} for d in subtraces}
    # deze gegevens vult de app nooit zelf in: de vergunningverlener beantwoordt ze altijd (afspraak Nico 1-11).
    # Wat de tekening erover zegt, tonen we als hint bij de vraag.
    hint = {}
    for naam, rij in basis.items():
        for veld in ALTIJD_VRAGEN:
            if rij.get(veld) not in (None, ""):
                hint[(naam, veld)] = rij.pop(veld)
    antw = st.session_state.setdefault("antwoorden", {})  # (trace, veld) -> waarde
    uit: set = st.session_state.setdefault("uit", set())  # uitgezette deeltracés
    uit &= set(sub_van)

    def vraag_van(veld):
        return info.get(veld, {}).get("vraag aan vergunningverlener") or veld

    def invoer_van(naam):
        rij = dict(basis[naam])
        for (tn, v), w in antw.items():
            if tn == naam and w not in (None, ""):
                rij[v] = w
        return rij

    def invulveld(veld, label, sleutel):
        v = info.get(veld, {})
        if veld in keuzes:
            return st.selectbox(label, ["(kies)"] + keuzes[veld], key=sleutel)
        eenheid = f" ({v.get('eenheid')})" if v.get("eenheid") else ""
        return st.text_input(label + eenheid, key=sleutel, placeholder="getal, bijv. 1.2")

    def bewaar(naam, veld, waarde):
        if waarde in (None, "", "(kies)"):
            return
        if veld not in keuzes:
            try:
                waarde = float(str(waarde).replace(",", "."))
            except ValueError:
                st.error(f"'{waarde}' is geen getal ({vraag_van(veld)}).")
                return
        antw[(naam, veld)] = waarde

    uitkomsten = {}
    for naam, d in sub_van.items():
        inv = invoer_van(naam)
        r = K.toets_deel(t, inv, d["toetszones"], d["geo"])
        open_ = [v for v in ALTIJD_VRAGEN if v in velden and inv.get(v) in (None, "") and v not in r.vragen]
        if open_ and r.zones:
            r.vragen = sorted(r.vragen + open_)
            if K.OORDELEN.index(r.oordeel) > K.OORDELEN.index("onvolledig"):
                r.oordeel = "onvolledig"
        uitkomsten[naam] = (r.oordeel, [(d, r)])
    aan = [n for n in uitkomsten if n not in uit]
    ernst = lambda n: (K.OORDELEN.index(uitkomsten[n][0]), int(n.split("#")[-1]))  # noqa: E731
    volgorde = sorted(uitkomsten, key=ernst)

    def kering_tekst(d):
        return ", ".join(d.get("keringen", [])) or "-"

    # ---------------------------------------------------------- 1. samenvatting
    st.caption(f"{len(lijnen)} lijn(en) in de upload, opgesplitst in {len(subtraces)} deeltracé(s) binnen de zones"
               + (f"; {len(uit)} uitgezet." if uit else "."))
    if not aan:
        st.warning("Alle deeltracés staan uit. Zet er minstens één aan in de tabel hieronder.")
    else:
        zwaarst = K.zwaarste(uitkomsten[n][0] for n in aan)
        bij = sorted([n for n in aan if uitkomsten[n][0] == zwaarst], key=ernst)
        kop = {
            "vergunning": f"Vergunningplichtig: {nrs_tekst(bij)}.",
            "maatwerk": f"Maatwerk nodig: {nrs_tekst(bij)} voldoet niet aan alle regels.",
            "onvolledig": f"Nog niet compleet: bij {'alle deeltracés' if len(bij) == len(aan) > 2 else nrs_tekst(bij)} "
                          f"ontbreken gegevens. Beantwoord de vragen hieronder.",
            "aandacht": f"Toegestaan, met aandachtspunten bij {nrs_tekst(bij)}.",
            "voorschriften": "Toegestaan met voorschriften.",
            "geen_specifieke_regels": "Geen specifieke regels van toepassing.",
            "buiten_beperkingengebied": "Buiten het beperkingengebied.",
        }[zwaarst]
        tellers = " ".join(f"<span style='border:1px solid {OORDEEL_HEX[o]};border-radius:12px;padding:2px 10px;"
                           f"margin-right:6px;white-space:nowrap'><b>{sum(uitkomsten[n][0] == o for n in aan)}</b> "
                           f"{OORDEEL_TEKST[o].lower()}</span>"
                           for o in K.OORDELEN if any(uitkomsten[n][0] == o for n in aan))
        st.markdown(f"<div style='border-left:6px solid {OORDEEL_HEX[zwaarst]};background:rgba(0,0,0,.03);"
                    f"padding:12px 16px;border-radius:4px;margin-bottom:8px'>"
                    f"<div style='font-size:1.25em;font-weight:600;margin-bottom:8px'>{kop}</div>{tellers}</div>",
                    unsafe_allow_html=True)

    # ---------------------------------------------------------- 2. kaart
    kaart(sub_van, uitkomsten, uit, lagen, zonenaam, keringnaam)

    # ---------------------------------------------------------- 3. overzicht met vinkjes
    st.subheader("Overzicht per deeltracé")
    st.caption("Haal het vinkje weg om een deeltracé niet mee te nemen; de samenvatting, de kaart en de vragen "
               "rekenen dan zonder dat stuk.")
    st.session_state["_volgorde"] = volgorde
    editor_key = f"overzicht_{casus}_{st.session_state.get('_editor', 0)}"

    def vinkjes_verwerken():
        wijz = st.session_state.get(editor_key, {}).get("edited_rows", {})
        for i, w in wijz.items():
            if "meenemen" in w:
                n = st.session_state["_volgorde"][int(i)]
                (st.session_state["uit"].discard if w["meenemen"] else st.session_state["uit"].add)(n)
        st.session_state["_editor"] = st.session_state.get("_editor", 0) + 1

    overzicht = pd.DataFrame([{
        "meenemen": n not in uit, "tracé": nr(n),
        "oordeel": f"{OORDEEL_STIP[uitkomsten[n][0]]} {OORDEEL_TEKST[uitkomsten[n][0]]}",
        "soort": SOORT_TEKST.get(sub_van[n]["soort"], sub_van[n]["soort"]),
        "lengte (m)": round(sub_van[n]["hoofdlijn"].length, 1),
        "zone": ", ".join(zone_kort(z) for z in sorted(sub_van[n]["toetszones"])) or "-",
        "kering": kering_tekst(sub_van[n]), "kruising": sub_van[n].get("wijze") or "-",
    } for n in volgorde])
    st.data_editor(overzicht, key=editor_key, on_change=vinkjes_verwerken, hide_index=True, width="stretch",
                   disabled=[c for c in overzicht.columns if c != "meenemen"],
                   column_config={"meenemen": st.column_config.CheckboxColumn("mee", width="small"),
                                  "tracé": st.column_config.TextColumn(width="small"),
                                  "lengte (m)": st.column_config.NumberColumn(format="%.1f", width="small")})

    # ---------------------------------------------------------- 4. vragen
    vragen = sorted({(n, v) for n in aan for _, r in uitkomsten[n][1] for v in r.vragen})
    if vragen:
        per_veld: dict[str, list[str]] = {}
        for naam, v in vragen:
            per_veld.setdefault(v, []).append(naam)
        st.subheader(f"Nog {len(per_veld)} gegeven{'s' if len(per_veld) > 1 else ''} nodig")
        st.caption("Deze gegevens staan niet in het bestand maar zijn nodig voor de toetsing. "
                   "Vul ze in en klik op 'Antwoorden opslaan'; daarna kunnen er nieuwe vragen volgen.")
        with st.form("vragen"):
            invul = []
            for veld, namen in per_veld.items():
                st.markdown(f"**{vraag_van(veld)}** <span style='color:grey'>· {'alle tracés' if len(namen) == len(aan) > 2 else 'tracé ' + nrs_tekst(sorted(namen, key=ernst))}</span>",
                            unsafe_allow_html=True)
                hints = {}
                for n in namen:
                    if (n, veld) in hint:
                        hints.setdefault(str(hint[(n, veld)]), []).append(n)
                if hints:
                    st.caption("Volgens de tekening: " + "; ".join(f"{w} bij {nrs_tekst(sorted(ns, key=ernst))}"
                                                                   for w, ns in hints.items()) + ". Bevestig of kies anders.")
                c1, c2 = st.columns([1, 2])
                with c1:
                    invul.append((veld, namen, invulveld(veld, "Voor alle genoemde tracés", f"q_{veld}_alle")))
                if len(namen) > 1:
                    with c2:
                        with st.expander("Per tracé verschillend"):
                            for n in namen:
                                invul.append((veld, [n], invulveld(veld, f"Tracé {nr(n)}", f"q_{veld}_{n}")))
            if st.form_submit_button("Antwoorden opslaan", type="primary"):
                for veld, namen, w in invul:
                    for n in namen:
                        bewaar(n, veld, w)
                st.rerun()
    elif aan:
        st.success("Alle gegevens die nodig zijn, zijn bekend.")

    with st.expander("Gegevens per tracé bekijken of aanpassen"):
        tabel = pd.DataFrame([{"tracé": nr(n), **invoer_van(n)} for n in sorted(sub_van, key=ernst)],
                             columns=["tracé", *velden])
        st.dataframe(tabel, width="stretch", hide_index=True)
        with st.form("aanpassen"):
            c1, c2, c3 = st.columns(3)
            tn = c1.selectbox("Tracé", sorted(sub_van, key=lambda n: int(n.split("#")[-1])), format_func=nr)
            vd = c2.selectbox("Gegeven", velden, format_func=vraag_van)
            wd = c3.text_input("Nieuwe waarde (bij keuzevelden: een van de opties)")
            if st.form_submit_button("Waarde opslaan"):
                if vd in keuzes and wd not in keuzes[vd]:
                    st.error(f"Kies een van: {', '.join(keuzes[vd])}")
                else:
                    bewaar(tn, vd, wd)
                    st.rerun()
        if antw and st.button("Alle antwoorden wissen"):
            antw.clear()
            st.rerun()

    # ---------------------------------------------------------- 5. per deeltracé
    st.subheader("Toelichting per deeltracé")
    regels_rap, voorschr_rap = [], []
    for naam in volgorde:
        oordeel, delen = uitkomsten[naam]
        d, r = delen[0]
        meegenomen = naam not in uit
        with st.container(border=True):
            ha = f" + {d['huisaansluitingen']} huisaansluiting(en)" if d.get("huisaansluitingen") else ""
            st.markdown(f"<div style='display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap'>"
                        f"<span style='font-size:1.1em'><b>{nr(naam)}</b> · {SOORT_TEKST.get(d['soort'], d['soort'])} · "
                        f"{d['hoofdlijn'].length:.0f} m{ha} · kering {kering_tekst(d)}</span>"
                        f"{badge(oordeel) if meegenomen else '<i>uitgezet, telt niet mee</i>'}</div>",
                        unsafe_allow_html=True)
            if meegenomen:
                st.markdown("**Waarom dit oordeel**  \n" + "  \n".join(waarom(oordeel, d, r, zonenaam, keringnaam,
                                                                                  regime_artikel)))
                if r.vragen:
                    st.markdown("**Wat ontbreekt**  \n" + "  \n".join(f"- {vraag_van(v)}" for v in r.vragen))
                vs = list(dict.fromkeys(x.tekst for x in r.regels if x.uitkomst in ("voldoet", "voorschrift") and x.tekst))
                if vs:
                    st.markdown("**Voorschriften**\n\n" + "\n".join(f"- {v}" for v in vs))
            with st.expander("Technische details"):
                g = d["geo"]
                hoek = "-" if g.get("hoek_met_kering") is None else f"{g['hoek_met_kering']:.0f}°"
                st.markdown(f"Bestand: {d['groep']} (bronlijnen {', '.join(map(str, d['bronnen']))}) · "
                            f"getoetst op: {', '.join(r.zones) or 'geen'} · geraakt: {', '.join(sorted(d['zones']))} · "
                            f"kruising: {d.get('wijze') or '-'} · regime: {', '.join(r.regimes) or '-'} · "
                            f"hoek met kering {hoek} · dwars-aandeel kernzone {g.get('dwars_aandeel_kernzone')} · "
                            f"kruin geraakt: {'ja' if g.get('raakt_kruin') else 'nee'}")
                if r.regels:
                    df = pd.DataFrame([{"soort regel": x.stap, "artikel": x.artikel,
                                        "uitkomst": UITKOMST_TEKST.get(x.uitkomst, x.uitkomst), "tekst": x.tekst,
                                        "ontbreekt": ", ".join(x.ontbrekend), "_v": UITKOMST_VOLGORDE.get(x.uitkomst, 9)}
                                       for x in r.regels]).sort_values("_v").drop(columns="_v")
                    st.dataframe(df, width="stretch", hide_index=True)
        for x in r.regels:
            regels_rap.append({"tracé": nr(naam), "meegenomen": "ja" if meegenomen else "nee", "oordeel tracé":
                               OORDEEL_TEKST[oordeel], "soort regel": x.stap, "artikel": x.artikel,
                               "uitkomst": UITKOMST_TEKST.get(x.uitkomst, x.uitkomst), "tekst": x.tekst,
                               "ontbreekt": ", ".join(x.ontbrekend)})
            if meegenomen and x.uitkomst in ("voldoet", "voorschrift") and x.tekst:
                voorschr_rap.append({"tracé": nr(naam), "artikel": x.artikel, "voorschrift": x.tekst})

    # ---------------------------------------------------------- 6. voorschriften verzameld
    if voorschr_rap:
        st.subheader("Voorschriften voor de reactie aan de aanvrager")
        st.caption("Alle voorschriften van de meegenomen deeltracés, zonder dubbelingen. Kopieer met het icoon rechtsboven.")
        per_tekst: dict[str, list[str]] = {}
        for v in voorschr_rap:
            per_tekst.setdefault(v["voorschrift"], [])
            if v["tracé"] not in per_tekst[v["voorschrift"]]:
                per_tekst[v["voorschrift"]].append(v["tracé"])
        alle_aan = len(per_tekst) and all(len(ts) == len(aan) for ts in per_tekst.values())
        tekst = "\n".join(f"- {v}" + ("" if alle_aan else f" (tracé {', '.join(ts)})") for v, ts in per_tekst.items())
        st.code(tekst, language=None, wrap_lines=True)

    # ---------------------------------------------------------- 7. download
    overzicht_x = overzicht.drop(columns="oordeel").assign(
        meenemen=lambda df: df["meenemen"].map({True: "ja", False: "nee"}),
        oordeel=[OORDEEL_TEKST[uitkomsten[n][0]] for n in volgorde],
        bestand=[sub_van[n]["groep"] for n in volgorde],
        **{"ontbrekende gegevens": ["; ".join(vraag_van(v) for v in uitkomsten[n][1][0][1].vragen) for n in volgorde]})
    st.download_button("Download toetsing (Excel)",
                       excel_rapport(overzicht_x, pd.DataFrame(regels_rap),
                                     pd.DataFrame(voorschr_rap, columns=["tracé", "artikel", "voorschrift"])),
                       "kabelchecker_toetsing.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.caption(f"Template: {tpl_pad.name} · HOEK_MIN {t.params['HOEK_MIN']:.0f}° · DWARS_AANDEEL {t.params['DWARS_AANDEEL']:.2f}")


if __name__ == "__main__":
    main()
