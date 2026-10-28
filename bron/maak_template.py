"""Bouwt regeltemplate_v4_AGV.xlsx uit regels_data.py.

Gebruik: python maak_template.py <uitvoer.xlsx>
"""
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import regels_data as D

KOP = PatternFill("solid", fgColor="1F4E79")
KOPFONT = Font(bold=True, color="FFFFFF")
SOORTKLEUR = {"toets": "E2EFDA", "voorschrift": "FFFFFF", "aandacht": "FFF2CC"}

INVOER = {v[0] for v in D.INVOERVELDEN}
BEGRIP = {b[0]: b[2] for b in D.BEGRIPPEN}
PARAM = {p[0] for p in D.PARAMETERS}
GEO = {b[0] for b in D.BEGRIPPEN if b[3].startswith("geometrie")}


def namen(expr):
    return set(re.findall(r"\b[a-z_][a-z0-9_]*\b", expr or "")) - {"in"}


def velden(expr, gezien=None):
    """Invoervelden en geo-feiten waar een expressie (via begrippen) van afhangt."""
    gezien = gezien or set()
    uit = set()
    for n in namen(expr):
        if n in gezien:
            continue
        gezien.add(n)
        if n in INVOER:
            uit.add(n)
        elif n in GEO:
            uit.add("geo:" + n)
        elif n in BEGRIP:
            uit |= velden(BEGRIP[n], gezien)
    return uit


def blad(wb, titel, koppen, rijen, breedtes=None):
    ws = wb.create_sheet(titel)
    ws.append(koppen)
    for c in ws[1]:
        c.fill, c.font = KOP, KOPFONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rijen:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for i, b in enumerate(breedtes or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.auto_filter.ref = ws.dimensions
    return ws


def tekst_voldoen(art, wsv):
    return f"Art. {art}: {wsv[0].upper()}{wsv[1:]}"


def tekst_niet(art, wsv):
    w = wsv.rstrip(";.").rstrip(" en")
    return (f"Voldoet niet aan art. {art} ({w}). Maatwerk nodig: vraag de aanvrager "
            f"waarom hiervan wordt afgeweken en welke maatregelen worden genomen.")


def main(uit):
    wb = Workbook()
    ws = wb.active
    ws.title = "LEESMIJ"
    regels = [
        "Regeltemplate v4 - Waterschapsverordening AGV, kabels en leidingen (art. 2.15 en 3.14 t/m 3.18)",
        "",
        "Hoe de app toetst (vaste volgorde, geen sprongen):",
        "0. Het tracé wordt opgeknipt in delen: bij elke zonegrens uit de legger en bij elke wissel van eigenschap.",
        "1. Feiten afleiden (tabblad Begrippen) uit invoer (Invoervelden), geometrie en Parameters.",
        "2. Vergunningplicht: elke rij van Vergunning_2.15 waarvan zone en voorwaarde kloppen, geeft 'vergunning'.",
        "3. Regimes: elke rij waarvan de voorwaarde klopt, is van toepassing (meerdere tegelijk kan, bijv. vervangen = 3.17 + aanleg).",
        "4. Zorgplicht: elke rij met een geldend regime, een geraakte zone (ja) en een kloppende 'van toepassing als' wordt beoordeeld.",
        "   soort toets: voldoet -> voorschrift opnemen; voldoet niet -> maatwerk; onbekend -> vraag aan vergunningverlener.",
        "   soort voorschrift: alleen voorschrift opnemen.   soort aandacht: signaal voor de vergunningverlener.",
        "",
        "Eindoordeel per deel (zwaar naar licht): vergunning > maatwerk > onvolledig > aandacht > voorschriften > geen specifieke regels > buiten beperkingengebied.",
        "Het tracé krijgt het zwaarste oordeel van zijn delen; het rapport toont per deel waarom.",
        "",
        "Schrijfwijze van expressies: veldnamen en begrippen, vergelijkingen (== != < <= > >=), 'in (\"a\", \"b\")', AND, OR, NOT, haakjes.",
        "Tekst tussen dubbele aanhalingstekens. Parameters in HOOFDLETTERS. Een ontbrekend gegeven maakt een uitkomst 'onbekend', nooit 'nee'.",
        "AND met een onwaar deel is onwaar, OR met een waar deel is waar, ook als de rest onbekend is.",
        "",
        "Afspraken Nico (28-10-2026): ontgravingsdiepte = onderkant buis (beton: wanddikte erbij); persleiding = druk > 1 bar;",
        "beschermingszone-deel dat niet aan 3.15.6 voldoet = maatwerk; blazen en proefsleuf vallen onder 3.18 (informerend);",
        "mantelbuis is een eigen type; marges (HOEK_MIN, DWARS_AANDEEL) zijn instelbaar; ontbrekende gegevens vraagt de app aan de vergunningverlener.",
    ]
    for r in regels:
        ws.append([r])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 150

    blad(wb, "Parameters", ["naam", "waarde", "eenheid", "toelichting"],
         D.PARAMETERS, [24, 10, 10, 90])
    blad(wb, "Zones", ["code", "keringtype", "zone", "omschrijving"], D.ZONES, [10, 16, 24, 70])
    ws = wb["Zones"]
    ws.append([])
    ws.append(["geolaag", "bron", "gebruik"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for g in D.GEO_LAGEN:
        ws.append(list(g))

    blad(wb, "Invoervelden",
         ["veld", "vraag aan vergunningverlener", "type", "waarden", "eenheid", "nodig bij", "bron op tekening"],
         D.INVOERVELDEN, [24, 50, 8, 50, 8, 30, 50])
    blad(wb, "Begrippen", ["naam", "definitie", "formule", "bron"], D.BEGRIPPEN, [24, 55, 70, 20])

    blad(wb, "Vergunning_2.15",
         ["id", "lid", *D.ZONE_CODES, "voorwaarde", "benodigde gegevens", "uitkomst bij waar", "wsv-tekst"],
         [(i, lid, *["ja" if z in zones else "nee" for z in D.ZONE_CODES], vw,
           ", ".join(sorted(velden(vw))), "vergunning", t)
          for i, lid, zones, vw, t in D.VERGUNNING],
         [10, 8] + [7] * len(D.ZONE_CODES) + [60, 40, 12, 70])

    blad(wb, "Regimes", ["code", "artikel", "voorwaarde", "omschrijving"], D.REGIMES, [10, 14, 80, 60])

    rijen = []
    for art, reg, zones, vta, toets, soort, wsv in D.ZORGPLICHT:
        rijen.append((
            art, reg, *["ja" if z in zones else "nee" for z in D.ZONE_CODES],
            vta, toets, soort,
            ", ".join(sorted(velden(vta) | velden(toets))),
            wsv,
            tekst_voldoen(art, wsv) if soort != "aandacht" else wsv,
            tekst_niet(art, wsv) if soort == "toets" else "",
            "",
        ))
    ws = blad(wb, "Zorgplicht",
              ["artikel", "regime", *D.ZONE_CODES, "van toepassing als", "toets", "soort",
               "benodigde gegevens", "wsv-tekst", "voorschrift (bij voldoen)", "tekst bij niet voldoen", "opmerking"],
              rijen, [11, 9] + [7] * len(D.ZONE_CODES) + [30, 40, 11, 30, 60, 60, 60, 30])
    soortkol = 2 + len(D.ZONE_CODES) + 3
    for row in ws.iter_rows(min_row=2):
        kleur = SOORTKLEUR.get(row[soortkol - 1].value)
        if kleur:
            for c in row:
                c.fill = PatternFill("solid", fgColor=kleur)
    dv = DataValidation(type="list", formula1='"ja,nee"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"C2:{get_column_letter(2 + len(D.ZONE_CODES))}{ws.max_row}")
    dv2 = DataValidation(type="list", formula1='"toets,voorschrift,aandacht"')
    ws.add_data_validation(dv2)
    L = get_column_letter(soortkol)
    dv2.add(f"{L}2:{L}{ws.max_row}")

    blad(wb, "Testgevallen",
         ["id", "omschrijving", "zone", "eigenschappen", "geo-feiten", "verwacht oordeel", "toelichting verwachting",
          "uitkomst rekenkern", "geraakte artikelen", "status"],
         [(i, o, z, "; ".join(f"{k}={v}" for k, v in e.items()), "; ".join(f"{k}={v}" for k, v in g.items()), v, t, "", "", "")
          for i, o, z, e, g, v, t in D.TESTGEVALLEN],
         [6, 50, 8, 60, 50, 16, 40, 16, 50, 10])

    blad(wb, "Metadata", ["sleutel", "waarde"], [
        ("template_versie", "4.0"),
        ("laatst_bijgewerkt", date.today().isoformat()),
        ("waterschap", "AGV (Amstel, Gooi en Vecht)"),
        ("bron", "Waterschapsverordening AGV, CVDR702443, art. 2.15 en 3.14-3.18"),
        ("opvolger_van", "regeltemplate_v3.xlsx"),
        ("aantal_zorgplichtregels", len(D.ZORGPLICHT)),
    ], [24, 70])

    wb.save(uit)
    print("geschreven:", uit, "| zorgplichtregels:", len(D.ZORGPLICHT))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "regeltemplate_v4_AGV.xlsx")
